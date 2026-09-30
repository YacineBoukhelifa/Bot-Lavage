import os
import sys
from datetime import datetime
from zoneinfo import ZoneInfo
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

os.environ.setdefault("BOT_TOKEN", "test-token")
os.environ.setdefault("WEBHOOK_SECRET", "test-secret")


@pytest.fixture()
def client(tmp_path, monkeypatch):
    from bot import app as app_module

    monkeypatch.setattr(app_module.config, "DB_PATH", str(tmp_path / "global5.sqlite3"))
    monkeypatch.setattr(app_module.config, "GROUPE_AUTORISE", 0)
    monkeypatch.setattr(app_module.config, "AUTORISATIONS", {"saisie": [], "lecture": "*"})
    monkeypatch.setattr(app_module.config, "EXCEL_LOCAL_DIR", str(tmp_path / "exports"))
    app_module.db.init_db()

    sent = []
    next_id = [1]

    def fake_send_message(chat_id, text, reply_markup=None, parse_mode=None):
        mid = next_id[0]
        next_id[0] += 1
        sent.append({
            "kind": "send", "chat_id": chat_id, "text": text,
            "reply_markup": reply_markup, "parse_mode": parse_mode, "mid": mid,
        })
        return {"ok": True, "result": {"message_id": mid}}

    def fake_send_photo(chat_id, photo_bytes, caption=None, reply_markup=None, parse_mode=None):
        mid = next_id[0]
        next_id[0] += 1
        sent.append({
            "kind": "photo", "chat_id": chat_id, "text": caption, "caption": caption,
            "size": len(photo_bytes), "reply_markup": reply_markup, "parse_mode": parse_mode, "mid": mid,
        })
        return {"ok": True, "result": {"message_id": mid}}

    def fake_send_document(chat_id, file_path, caption=None):
        mid = next_id[0]
        next_id[0] += 1
        sent.append({
            "kind": "document", "chat_id": chat_id, "text": caption, "caption": caption,
            "path": file_path, "mid": mid,
        })
        return {"ok": True, "result": {"message_id": mid}}

    def fake_edit_message_text(chat_id, message_id, text, reply_markup=None, parse_mode=None):
        sent.append({
            "kind": "edit", "chat_id": chat_id, "mid": message_id, "text": text,
            "reply_markup": reply_markup, "parse_mode": parse_mode,
        })
        return {"ok": True, "result": {"message_id": message_id}}

    def fake_answer_cbq(callback_query_id, text=None):
        sent.append({"kind": "answer_cbq", "id": callback_query_id})

    monkeypatch.setattr(app_module.telegram_client, "send_message", fake_send_message)
    monkeypatch.setattr(app_module.telegram_client, "send_photo", fake_send_photo)
    monkeypatch.setattr(app_module.telegram_client, "send_document", fake_send_document)
    monkeypatch.setattr(app_module.telegram_client, "edit_message_text", fake_edit_message_text)
    monkeypatch.setattr(app_module.telegram_client, "answer_callback_query", fake_answer_cbq)

    test_client = app_module.app.test_client()
    test_client.sent = sent
    test_client.app_module = app_module
    return test_client


def _post_msg(client, text, hh=None, mm=None, reply_to_mid=None, date=None):
    config = client.app_module.config
    tz = ZoneInfo(config.TIMEZONE)
    n = datetime.now(tz)
    if date is None:
        date = n.strftime("%Y-%m-%d")
    if hh is None:
        hh = n.hour
    if mm is None:
        mm = n.minute
    y, m, d = (int(x) for x in date.split("-"))
    ts = int(datetime(y, m, d, hh, mm, tzinfo=tz).timestamp())
    payload = {
        "message": {
            "chat": {"id": 100, "type": "group"},
            "from": {"id": 1, "first_name": "Manager"},
            "text": text,
            "date": ts,
        }
    }
    if reply_to_mid is not None:
        payload["message"]["reply_to_message"] = {"message_id": reply_to_mid, "from": {"is_bot": True}}
    return client.post(f"/webhook/{config.WEBHOOK_SECRET}", json=payload)


def _post_cb(client, data, message_id):
    config = client.app_module.config
    return client.post(
        f"/webhook/{config.WEBHOOK_SECRET}",
        json={"callback_query": {
            "id": "cb1", "data": data, "from": {"id": 1, "first_name": "Manager"},
            "message": {"chat": {"id": 100, "type": "group"}, "message_id": message_id},
        }},
    )


def _last(sent):
    for e in reversed(sent):
        if e["kind"] != "answer_cbq":
            return e
    return None


def test_global_full_workflow_5_lines(client):
    """
    Scenario global complet avec les 5 lignes :
    1. Demarrage du shift avec 5 objectifs horaires
    2. Saisie guidee : point injecte (5 valeurs) + validation
    3. Commandes directes individuelles /l4 et /l5
    4. Commande multi-lignes /prod a=... s=... l3=... l4=... l5=...
    5. Questionnaire dynamique de pause dejeuner complet pour les 5 lignes
    6. Verifications /recap, /graph (photo) et /export (document xlsx)
    7. Cloture du shift avec /fin et synthese globale
    """
    tz = ZoneInfo(client.app_module.config.TIMEZONE)
    n = datetime.now(tz)
    today = n.strftime("%Y-%m-%d")

    from bot import db as db_module, logic as logic_module

    # Neutraliser cloture opportuniste au demarrage
    conn = db_module.get_connection()
    conn.execute(
        "INSERT OR REPLACE INTO syntheses_envoyees (date, poste, envoyee_a) VALUES (?, 1, ?)",
        (today, datetime.now().isoformat()),
    )
    conn.commit()
    conn.close()

    # 1. Demarrage du shift
    _post_msg(client, "/start_day")
    prompt_obj = _last(client.sent)
    assert "Objectifs horaires" in prompt_obj["text"]

    # Reponse avec 5 valeurs pour les 5 lignes : Auto, Semi-auto, L03, L04, L05
    _post_msg(client, "133 160 80 80 80", reply_to_mid=prompt_obj["mid"])
    res_start = _last(client.sent)
    assert "démarré" in res_start["text"].lower()

    # 2. Saisie guidee au point 09:00
    conn = db_module.get_connection()
    logic_module.set_interaction_state(
        conn, "100", "1", "ATTENTE_CUMULS",
        {"poste": 1, "date": today, "heure": "09:00"}, n, 42,
    )
    conn.close()

    client.sent.clear()
    _post_msg(client, "133 160 80 80 80", reply_to_mid=42)
    card09 = _last(client.sent)
    assert "à valider" in card09["text"]

    # Validation
    client.sent.clear()
    _post_cb(client, "guide_valider", message_id=card09["mid"])
    report09 = _last(client.sent)
    assert report09["kind"] == "send"
    assert "RAPPORT HORAIRE" in report09["text"]
    for nom in ["Ligne Auto", "Ligne Semi-auto", "Ligne 03", "Ligne 04", "Ligne 05"]:
        assert nom in report09["text"]

    # 3. Commandes directes individuelles /l4 et /l5 au point 10:00 (simulate hh=10, mm=0)
    # Ligne 4 et Ligne 5 utilisent la date du shift actif
    client.sent.clear()
    _post_msg(client, "/l4 160", hh=10, mm=0)
    rep_l4 = _last(client.sent)
    assert "Ligne 04" in rep_l4["text"]
    assert "Production : 80 pcs" in rep_l4["text"]  # 160 - 80 = 80

    _post_msg(client, "/l5 160", hh=10, mm=0)
    rep_l5 = _last(client.sent)
    assert "Ligne 05" in rep_l5["text"]
    assert "Production : 80 pcs" in rep_l5["text"]

    # 4. Commande /prod avec toutes les 5 lignes
    client.sent.clear()
    _post_msg(client, "/prod a=399 s=480 l3=240 l4=240 l5=240", hh=11, mm=0)
    rep_prod = _last(client.sent)
    assert "RAPPORT HORAIRE" in rep_prod["text"]
    for nom in ["Ligne Auto", "Ligne Semi-auto", "Ligne 03", "Ligne 04", "Ligne 05"]:
        assert nom in rep_prod["text"]

    # 5. Questionnaire dynamique pause dejeuner pour les 5 lignes
    conn = db_module.get_connection()
    logic_module.set_interaction_state(
        conn, "100", "1", "ATTENTE_PAUSE_DEJEUNER",
        {"date": today, "poste": 1, "codes_restants": ["A", "S", "SKD", "L4", "L5"]}, n, 1001,
    )
    conn.close()

    # Ligne Auto : Oui (pause 12:00)
    client.sent.clear()
    _post_cb(client, "pause_dej|oui", message_id=1001)
    q2 = _last(client.sent)
    assert "Ligne Semi-auto" in q2["text"]

    # Ligne Semi-auto : Non (pause a 13:00)
    _post_cb(client, "pause_dej|non", message_id=q2["mid"])
    q3 = _last(client.sent)
    assert "Ligne 03" in q3["text"]

    # Ligne 03 : Non
    _post_cb(client, "pause_dej|non", message_id=q3["mid"])
    q4 = _last(client.sent)
    assert "Ligne 04" in q4["text"]

    # Ligne 04 : Oui (pause 12:00)
    _post_cb(client, "pause_dej|oui", message_id=q4["mid"])
    q5 = _last(client.sent)
    assert "Ligne 05" in q5["text"]

    # Ligne 05 : Non
    _post_cb(client, "pause_dej|non", message_id=q5["mid"])
    p12 = _last(client.sent)
    assert "Point de contrôle 12:00" in p12["text"]
    assert p12["reply_markup"] == {"force_reply": True, "selective": True}

    # Verification en base de la pause dejeuner des 5 lignes
    conn = db_module.get_connection()
    assert logic_module.get_pause_dejeuner(conn, today, "A") == "12:00"
    assert logic_module.get_pause_dejeuner(conn, today, "S") == "13:00"
    assert logic_module.get_pause_dejeuner(conn, today, "SKD") == "13:00"
    assert logic_module.get_pause_dejeuner(conn, today, "L4") == "12:00"
    assert logic_module.get_pause_dejeuner(conn, today, "L5") == "13:00"
    conn.close()

    # Reponse au point 12:00 avec 5 cumuls
    _post_msg(client, "465 640 320 280 320", reply_to_mid=p12["mid"])
    card12 = _last(client.sent)
    client.sent.clear()
    _post_cb(client, "guide_valider", message_id=card12["mid"])
    rep12 = _last(client.sent)
    assert "RAPPORT HORAIRE" in rep12["text"]

    # 6. Test /recap
    client.sent.clear()
    _post_msg(client, "/recap")
    recap = _last(client.sent)
    assert "RECAP" in recap["text"]
    for nom in ["Ligne Auto", "Ligne Semi-auto", "Ligne 03", "Ligne 04", "Ligne 05"]:
        assert nom in recap["text"]

    # Test /graph
    client.sent.clear()
    _post_msg(client, "/graph")
    photos = [m for m in client.sent if m["kind"] == "photo"]
    assert len(photos) == 1
    assert photos[0]["size"] > 1000

    # Test /export (Excel)
    client.sent.clear()
    _post_msg(client, "/export")
    docs = [m for m in client.sent if m["kind"] == "document"]
    assert len(docs) == 1
    assert os.path.exists(docs[0]["path"])

    # Verifier le classeur Excel genere
    from openpyxl import load_workbook
    wb = load_workbook(docs[0]["path"])
    assert "Donnees" in wb.sheetnames
    assert "Synthese_Quotidienne" in wb.sheetnames
    assert "Vue_Mensuelle" in wb.sheetnames

    # Verifier que les 5 lignes (A, S, SKD, L4, L5) sont bien dans l'onglet Donnees
    ws_donnees = wb["Donnees"]
    lignes_codes_excel = {ws_donnees.cell(row=r, column=5).value for r in range(3, ws_donnees.max_row + 1)}
    assert {"A", "S", "SKD", "L4", "L5"}.issubset(lignes_codes_excel)

    # 7. Cloture du poste avec /fin
    conn = db_module.get_connection()
    conn.execute("DELETE FROM syntheses_envoyees WHERE date=? AND poste=1", (today,))
    conn.commit()
    conn.close()

    client.sent.clear()
    _post_msg(client, "/fin")
    fin_entries = client.sent
    photos_fin = [m for m in fin_entries if m["kind"] == "photo"]
    assert len(photos_fin) >= 1  # Graphique de fin envoye
    text_fin = [m["text"] for m in fin_entries if m["kind"] == "send"]
    assert any("SYNTHÈSE" in (t or "").upper() for t in text_fin)
